from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl

from fusion_clean import (
    MOTIF_TARGET_4,
    OMEGA_TARGET,
    _compute_step,
    nuclear_fusion_coarse_grain,
    sovereign_dynamics_step,
)
from fusion_clean import neutron_coarse_grain as neutron_coarse_grain_step


ROOT = Path(__file__).resolve().parent
IN_XLSX = ROOT / "fusionequation.xlsx"
OUT_XLSX = ROOT / "analysis_results" / "fusionequation_filled.xlsx"
OUT_XLSX_FALLBACK = ROOT / "analysis_results" / "fusionequation_filled_unlocked.xlsx"
OUT_JSON = ROOT / "analysis_results" / "fusionequation_score.json"
OUT_MD = ROOT / "analysis_results" / "fusionequation_score.md"


PHASE1_COL = 4  # D
PHASE2_COL = 5  # E
HYST_COL = 6  # F


def _norm_label(x: object) -> str:
    s = "" if x is None else str(x)
    return " ".join(s.strip().lower().split())


def _to_control_state(x: object) -> str:
    """
    Spreadsheet -> operator state.
    - Actively On  -> on
    - Actively Off -> off
    - No control   -> no_control
    Empty/unknown  -> no_control
    """
    if x is None:
        return "no_control"
    s = " ".join(str(x).strip().lower().split())
    if not s:
        return "no_control"
    if "no" in s and "control" in s:
        return "no_control"
    if s in ("no_control", "no-control"):
        return "no_control"
    if "actively" in s and "on" in s:
        return "on"
    if "actively" in s and "off" in s:
        return "off"
    if s == "on":
        return "on"
    if s == "off":
        return "off"
    return "no_control"


def _to_sheet_value(state: str) -> str:
    if state == "on":
        return "Actively On"
    if state == "off":
        return "Actively Off"
    return "No control"


ROW_TO_KEY = {
    "gdh": "gdh_gluon",
    "women's gaba-b": "female_gaba_b_latdorsi",
    "acetyl coa": "left_acetyl_coa",
    "left serotonin": "male_left_5ht",
    "left noradrenaline": "female_left_noradrenaline",
    "right alpha2": "right_alpha2",
    "5ht1a": "left_temporalis_5ht1a",
    "left estrogen": "left_estrogen",
    "right love": "right_love",
    "hypoxia": "hypoxia",
    "right dopamine": "right_dopamine",
    "big woman's vasopressin": "vasopressin_female",
    "small man's oxytocin": "male_oxytocin",
    "blood a muscle": "muscle_a",
    "blood b muscle": "muscle_b",
    "5ht1b": "right_5ht1b_synchrotron",
    "right androgen": "right_androgen",
    "left endorphin": "left_endorphin",
    "left d2": "left_frontalis_d2",
    "gaba-a": "right_occipitalis_gaba_a",
    "right acetylcholine": "right_acetylcholine",
    "left extraversion": "left_extraversion",
    "glucocorticoid": "glucocorticoid",
    "right cortisol": "right_cortisol",
}


@dataclass(frozen=True)
class PhaseProfiles:
    phase1: dict[str, str]
    phase2: dict[str, str]
    hysteresis: dict[str, str]
    sheet_specified: int
    sheet_total: int
    filled_output_path: str


def load_and_fill() -> PhaseProfiles:
    if not IN_XLSX.exists():
        raise FileNotFoundError(str(IN_XLSX))
    wb = openpyxl.load_workbook(IN_XLSX)
    ws = wb.active

    specified = 0
    total = 0
    p1: dict[str, str] = {}
    p2: dict[str, str] = {}
    hy: dict[str, str] = {}

    # Fill blanks with "No control" and build control dicts.
    for r in range(2, ws.max_row + 1):
        label = _norm_label(ws.cell(r, 1).value)
        if not label:
            continue
        key = ROW_TO_KEY.get(label)
        if not key:
            continue

        raw1 = ws.cell(r, PHASE1_COL).value
        raw2 = ws.cell(r, PHASE2_COL).value
        raw3 = ws.cell(r, HYST_COL).value
        for raw in (raw1, raw2, raw3):
            total += 1
            if raw is not None and str(raw).strip() != "":
                specified += 1

        s1 = _to_control_state(raw1)
        s2 = _to_control_state(raw2)
        s3 = _to_control_state(raw3)

        # If user left blank, write explicit "No control" so the sheet is complete.
        if raw1 is None or str(raw1).strip() == "":
            ws.cell(r, PHASE1_COL).value = _to_sheet_value(s1)
        if raw2 is None or str(raw2).strip() == "":
            ws.cell(r, PHASE2_COL).value = _to_sheet_value(s2)
        if raw3 is None or str(raw3).strip() == "":
            ws.cell(r, HYST_COL).value = _to_sheet_value(s3)

        p1[key] = s1
        p2[key] = s2
        hy[key] = s3

    # Ensure RIGHT_ALPHA2 exists as a separate row (the input sheet currently only has "left noradrenaline").
    # This is written only to the filled output workbook, not the user's original file.
    if "right_alpha2" not in p1:
        r = ws.max_row + 1
        ws.cell(r, 1).value = "right alpha2"
        ws.cell(r, PHASE1_COL).value = _to_sheet_value("no_control")
        ws.cell(r, PHASE2_COL).value = _to_sheet_value("no_control")
        # User intent: right_alpha2 is idle in day phases, but actively OFF in hysteresis (disinhibition window).
        ws.cell(r, HYST_COL).value = _to_sheet_value("off")
        p1["right_alpha2"] = "no_control"
        p2["right_alpha2"] = "no_control"
        hy["right_alpha2"] = "off"
        total += 3

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    filled_path = OUT_XLSX
    try:
        wb.save(OUT_XLSX)
    except PermissionError:
        # If the file is open in Excel, save to a different filename so the run remains reproducible.
        wb.save(OUT_XLSX_FALLBACK)
        filled_path = OUT_XLSX_FALLBACK
    return PhaseProfiles(
        phase1=p1,
        phase2=p2,
        hysteresis=hy,
        sheet_specified=specified,
        sheet_total=total,
        filled_output_path=str(filled_path),
    )


def minute_to_hhmm(minute: int) -> str:
    minute = int(minute) % (24 * 60)
    hh = minute // 60
    mm = minute % 60
    return f"{hh:02d}:{mm:02d}"


def is_hysteresis(minute: int) -> bool:
    return 90 <= int(minute) <= 180  # 01:30??3:00


def run_score(profiles: PhaseProfiles) -> dict[str, object]:
    # 16 windows x 90 minutes. Evaluate at the mid-point of each window.
    windows = []
    for i in range(16):
        start = i * 90
        mid = start + 45
        windows.append((i, start, mid))

    state = np.asarray(MOTIF_TARGET_4, dtype=float)
    records = []

    for i, start, mid in windows:
        hhmm = minute_to_hhmm(mid)
        if is_hysteresis(mid):
            regime = "hysteresis"
            control = dict(profiles.hysteresis)
        else:
            regime = "phase1" if (i % 2 == 0) else "phase2"
            control = dict(profiles.phase1 if regime == "phase1" else profiles.phase2)

        # Efficiency signals (0.8009 target) live in the base operator output.
        base = _compute_step(
            state,
            phase_fill=0.5,
            clock_hhmm=hhmm,
            edge15_gain=0.0,
            control_override=control,
        )
        ts = dict(base.get("tunnel_transfer_split") or {})
        transfer_eff = float(ts.get("transfer_efficiency", 0.0))
        static_eff = float(ts.get("static_efficiency", 0.0))

        out = sovereign_dynamics_step(
            state,
            phase_fill=0.5,
            clock_hhmm=hhmm,
            dt=1.0,
            control_override=control,
        )
        state = np.asarray(out["state_next"], dtype=float)
        omega = float(out.get("omega_next", float(np.linalg.norm(state))))
        omega_err = float(omega - OMEGA_TARGET)
        dist_target = float(np.linalg.norm(state - np.asarray(MOTIF_TARGET_4, dtype=float)))

        fusion = nuclear_fusion_coarse_grain(
            state,
            phase_fill=0.5,
            clock_hhmm=hhmm,
            control_override=control,
        )
        fusion_rate = float(fusion.get("fusion_rate", 0.0))
        fcomp = dict(fusion.get("components") or {})
        break_frac = float(fcomp.get("string_break_fraction", 0.0))
        z_proxy = float(fcomp.get("z_boson_proxy", 0.0))
        sigma_eff = float(fcomp.get("effective_string_tension", 0.0))

        neutron_obs = float(
            neutron_coarse_grain_step(
                state,
                phase_fill=0.5,
                clock_hhmm=hhmm,
                control_override=control,
            ).get("observable", 0.0)
        )
        neutrino_seed = float(abs(state[2]))

        records.append(
            {
                "window_index": i,
                "window_start_min": start,
                "window_mid_min": mid,
                "hhmm": hhmm,
                "regime": regime,
                "omega": omega,
                "omega_err": omega_err,
                "dist_to_motif_target": dist_target,
                "fusion_rate": fusion_rate,
                "break_frac": break_frac,
                "z_proxy": z_proxy,
                "sigma_eff": sigma_eff,
                "transfer_eff": transfer_eff,
                "static_eff": static_eff,
                "spark": float(out.get("spark", 0.0)),
                "neutron_star_obs": neutron_obs,
                "neutrino_seed": neutrino_seed,
            }
        )

    omega_errs = np.array([r["omega_err"] for r in records], dtype=float)
    dists = np.array([r["dist_to_motif_target"] for r in records], dtype=float)
    fusion_rates = np.array([r["fusion_rate"] for r in records], dtype=float)
    transfer_effs = np.array([r["transfer_eff"] for r in records], dtype=float)
    static_effs = np.array([r["static_eff"] for r in records], dtype=float)

    summary = {
        "rmse_omega": float(np.sqrt(float(np.mean(omega_errs**2)))),
        "mae_omega": float(np.mean(np.abs(omega_errs))),
        "max_abs_omega": float(np.max(np.abs(omega_errs))),
        "mean_dist_to_motif_target": float(np.mean(dists)),
        "max_dist_to_motif_target": float(np.max(dists)),
        "mean_fusion_rate": float(np.mean(fusion_rates)),
        "max_fusion_rate": float(np.max(fusion_rates)),
        "min_fusion_rate": float(np.min(fusion_rates)),
        "mean_transfer_eff": float(np.mean(transfer_effs)),
        "mean_static_eff": float(np.mean(static_effs)),
    }

    return {"summary": summary, "records": records}


def run_score_base_only(profiles: PhaseProfiles) -> dict[str, object]:
    """
    Score using ONLY the base step (no sovereign engineering/intent/dissip layers).
    This answers: does the sheet schedule itself keep omega near 7.4?
    """
    windows = []
    for i in range(16):
        start = i * 90
        mid = start + 45
        windows.append((i, start, mid))

    state = np.asarray(MOTIF_TARGET_4, dtype=float)
    records = []
    for i, start, mid in windows:
        hhmm = minute_to_hhmm(mid)
        if is_hysteresis(mid):
            regime = "hysteresis"
            control = dict(profiles.hysteresis)
        else:
            regime = "phase1" if (i % 2 == 0) else "phase2"
            control = dict(profiles.phase1 if regime == "phase1" else profiles.phase2)

        base = _compute_step(
            state,
            phase_fill=0.5,
            clock_hhmm=hhmm,
            edge15_gain=0.0,
            control_override=control,
        )
        ts = dict(base.get("tunnel_transfer_split") or {})
        state = np.asarray(base["state_out"], dtype=float)
        omega = float(np.linalg.norm(state))
        omega_err = float(omega - OMEGA_TARGET)
        dist_target = float(np.linalg.norm(state - np.asarray(MOTIF_TARGET_4, dtype=float)))

        fusion = nuclear_fusion_coarse_grain(
            state,
            phase_fill=0.5,
            clock_hhmm=hhmm,
            control_override=control,
        )
        fusion_rate = float(fusion.get("fusion_rate", 0.0))
        fcomp = dict(fusion.get("components") or {})

        records.append(
            {
                "window_index": i,
                "window_start_min": start,
                "window_mid_min": mid,
                "hhmm": hhmm,
                "regime": regime,
                "omega": omega,
                "omega_err": omega_err,
                "dist_to_motif_target": dist_target,
                "fusion_rate": fusion_rate,
                "break_frac": float(fcomp.get("string_break_fraction", 0.0)),
                "z_proxy": float(fcomp.get("z_boson_proxy", 0.0)),
                "sigma_eff": float(fcomp.get("effective_string_tension", 0.0)),
                "transfer_eff": float(ts.get("transfer_efficiency", 0.0)),
                "static_eff": float(ts.get("static_efficiency", 0.0)),
            }
        )

    omega_errs = np.array([r["omega_err"] for r in records], dtype=float)
    dists = np.array([r["dist_to_motif_target"] for r in records], dtype=float)
    fusion_rates = np.array([r["fusion_rate"] for r in records], dtype=float)
    transfer_effs = np.array([r["transfer_eff"] for r in records], dtype=float)
    static_effs = np.array([r["static_eff"] for r in records], dtype=float)
    summary = {
        "rmse_omega": float(np.sqrt(float(np.mean(omega_errs**2)))),
        "mae_omega": float(np.mean(np.abs(omega_errs))),
        "max_abs_omega": float(np.max(np.abs(omega_errs))),
        "mean_dist_to_motif_target": float(np.mean(dists)),
        "max_dist_to_motif_target": float(np.max(dists)),
        "mean_fusion_rate": float(np.mean(fusion_rates)),
        "max_fusion_rate": float(np.max(fusion_rates)),
        "min_fusion_rate": float(np.min(fusion_rates)),
        "mean_transfer_eff": float(np.mean(transfer_effs)),
        "mean_static_eff": float(np.mean(static_effs)),
    }
    return {"summary": summary, "records": records}

def write_report(profiles: PhaseProfiles, scored: dict[str, object]) -> None:
    OUT_JSON.write_text(json.dumps(scored, ensure_ascii=False, indent=2), encoding="utf-8")

    s = dict(scored.get("summary") or {})
    # Legacy "percent closure" used in prior notes: 1 - RMSE(omega)/OMEGA_TARGET.
    rmse_omega = float(s.get("rmse_omega") or 0.0)
    closure_percent = float(100.0 * (1.0 - rmse_omega / max(OMEGA_TARGET, 1.0e-12)))
    specified = int(getattr(profiles, "sheet_specified", 0))
    total = int(getattr(profiles, "sheet_total", 0))
    specified_ratio = float(specified) / float(total) if total > 0 else 0.0
    lines = [
        "# Fusion Equation Fill + Closure Score",
        "",
        f"- input: `{IN_XLSX.name}`",
        f"- filled_output: `{OUT_XLSX.name}`",
        "",
        "## Sheet Spec Coverage",
        f"- specified_cells: `{specified}` / `{total}` ({specified_ratio:.2%})",
        "",
        "## Closure (Omega)",
        f"- rmse_omega: `{s.get('rmse_omega')}`",
        f"- closure_percent_legacy: `{closure_percent:.4f}`  (100*(1 - rmse_omega/7.4))",
        f"- mae_omega: `{s.get('mae_omega')}`",
        f"- max_abs_omega: `{s.get('max_abs_omega')}`",
        "",
        "## Motif Target Distance",
        f"- mean_dist_to_motif_target: `{s.get('mean_dist_to_motif_target')}`",
        f"- max_dist_to_motif_target: `{s.get('max_dist_to_motif_target')}`",
        "",
        "## Fusion Proxy",
        f"- mean_fusion_rate: `{s.get('mean_fusion_rate')}`",
        f"- max_fusion_rate: `{s.get('max_fusion_rate')}`",
        f"- min_fusion_rate: `{s.get('min_fusion_rate')}`",
        "",
        "## Window Records (16 x 90min, evaluated at midpoints)",
        "",
        "|i|mid|regime|spark|break|Z|sigma_eff|omega_err|dist_target|fusion_rate|day_neutron|hyst_neutrino|",
        "|-:|:--:|:--:|--:|--:|--:|--:|--:|--:|--:|--:|--:|",
    ]
    for r in scored.get("records") or []:
        lines.append(
            "|{i}|{hhmm}|{reg}|{sp:.6g}|{bf:.6g}|{zp:.6g}|{se:.6g}|{oe:.6g}|{dt:.6g}|{fr:.6g}|{no:.6g}|{ns:.6g}|".format(
                i=int(r["window_index"]),
                hhmm=str(r["hhmm"]),
                reg=str(r["regime"]),
                sp=float(r.get("spark", 0.0)),
                bf=float(r.get("break_frac", 0.0)),
                zp=float(r.get("z_proxy", 0.0)),
                se=float(r.get("sigma_eff", 0.0)),
                oe=float(r["omega_err"]),
                dt=float(r["dist_to_motif_target"]),
                fr=float(r["fusion_rate"]),
                no=float(r["neutron_star_obs"]),
                ns=float(r["neutrino_seed"]),
            )
        )
    lines.append("")

    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    profiles = load_and_fill()
    scored = run_score(profiles)
    base_only = run_score_base_only(profiles)

    # Bundle both into one json.
    bundle = {"sovereign": scored, "base_only": base_only}
    OUT_JSON.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")

    # Write md with both summaries.
    s = dict(scored.get("summary") or {})
    b = dict(base_only.get("summary") or {})
    # Legacy "percent closure": 100 * (1 - RMSE(omega)/OMEGA_TARGET).
    s_rmse = float(s.get("rmse_omega") or 0.0)
    b_rmse = float(b.get("rmse_omega") or 0.0)
    s_pct = float(100.0 * (1.0 - s_rmse / max(OMEGA_TARGET, 1.0e-12)))
    b_pct = float(100.0 * (1.0 - b_rmse / max(OMEGA_TARGET, 1.0e-12)))
    specified = int(getattr(profiles, "sheet_specified", 0))
    total = int(getattr(profiles, "sheet_total", 0))
    specified_ratio = float(specified) / float(total) if total > 0 else 0.0
    lines = [
        "# Fusion Equation Fill + Closure Score",
        "",
        f"- input: `{IN_XLSX.name}`",
        f"- filled_output: `{getattr(profiles, 'filled_output_path', str(OUT_XLSX))}`",
        "",
        "## Sheet Spec Coverage",
        f"- specified_cells: `{specified}` / `{total}` ({specified_ratio:.2%})",
        "",
        "## Closure (Omega) - sovereign_dynamics_step (has built-in restore/intent/dissip)",
        f"- rmse_omega: `{s.get('rmse_omega')}`",
        f"- closure_percent_legacy: `{s_pct:.4f}`  (100*(1 - rmse_omega/7.4))",
        f"- mae_omega: `{s.get('mae_omega')}`",
        f"- max_abs_omega: `{s.get('max_abs_omega')}`",
        "",
        "## Closure (Omega) - base_only (_compute_step only; no sovereign closure controller)",
        f"- rmse_omega: `{b.get('rmse_omega')}`",
        f"- closure_percent_legacy: `{b_pct:.4f}`  (100*(1 - rmse_omega/7.4))",
        f"- mae_omega: `{b.get('mae_omega')}`",
        f"- max_abs_omega: `{b.get('max_abs_omega')}`",
        "",
        "## Motif Target Distance (sovereign / base_only)",
        f"- mean_dist_to_motif_target: `{s.get('mean_dist_to_motif_target')}` / `{b.get('mean_dist_to_motif_target')}`",
        f"- max_dist_to_motif_target: `{s.get('max_dist_to_motif_target')}` / `{b.get('max_dist_to_motif_target')}`",
        "",
        "## Efficiency (0.8009) (sovereign / base_only)",
        f"- mean_transfer_eff: `{s.get('mean_transfer_eff')}` / `{b.get('mean_transfer_eff')}`",
        f"- mean_static_eff: `{s.get('mean_static_eff')}` / `{b.get('mean_static_eff')}`",
        "",
        "## Fusion Proxy (sovereign / base_only)",
        f"- mean_fusion_rate: `{s.get('mean_fusion_rate')}` / `{b.get('mean_fusion_rate')}`",
        f"- max_fusion_rate: `{s.get('max_fusion_rate')}` / `{b.get('max_fusion_rate')}`",
        f"- min_fusion_rate: `{s.get('min_fusion_rate')}` / `{b.get('min_fusion_rate')}`",
        "",
        "## Window Records (sovereign)",
        "",
        "|i|mid|regime|spark|break|Z|sigma_eff|omega_err|dist_target|fusion_rate|day_neutron|hyst_neutrino|",
        "|-:|:--:|:--:|--:|--:|--:|--:|--:|--:|--:|--:|--:|",
    ]
    for r in scored.get("records") or []:
        lines.append(
            "|{i}|{hhmm}|{reg}|{sp:.6g}|{bf:.6g}|{zp:.6g}|{se:.6g}|{oe:.6g}|{dt:.6g}|{fr:.6g}|{no:.6g}|{ns:.6g}|".format(
                i=int(r["window_index"]),
                hhmm=str(r["hhmm"]),
                reg=str(r["regime"]),
                sp=float(r.get("spark", 0.0)),
                bf=float(r.get("break_frac", 0.0)),
                zp=float(r.get("z_proxy", 0.0)),
                se=float(r.get("sigma_eff", 0.0)),
                oe=float(r["omega_err"]),
                dt=float(r["dist_to_motif_target"]),
                fr=float(r["fusion_rate"]),
                no=float(r["neutron_star_obs"]),
                ns=float(r["neutrino_seed"]),
            )
        )
    lines.append("")
    lines.append("## Window Records (base_only)")
    lines.append("")
    lines.append("|i|mid|regime|break|Z|sigma_eff|omega_err|dist_target|fusion_rate|")
    lines.append("|-:|:--:|:--:|--:|--:|--:|--:|--:|--:|")
    for r in base_only.get("records") or []:
        lines.append(
            "|{i}|{hhmm}|{reg}|{bf:.6g}|{zp:.6g}|{se:.6g}|{oe:.6g}|{dt:.6g}|{fr:.6g}|".format(
                i=int(r["window_index"]),
                hhmm=str(r["hhmm"]),
                reg=str(r["regime"]),
                bf=float(r.get("break_frac", 0.0)),
                zp=float(r.get("z_proxy", 0.0)),
                se=float(r.get("sigma_eff", 0.0)),
                oe=float(r["omega_err"]),
                dt=float(r["dist_to_motif_target"]),
                fr=float(r["fusion_rate"]),
            )
        )
    lines.append("")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(str(OUT_MD))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
